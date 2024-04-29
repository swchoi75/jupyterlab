import csv
import openpyxl as op
from pathlib import Path


# Path
path = Path(__file__).parent.parent


# Functions


def main():
    # Filenames
    input_file = path / "data" / "MYP consolidation template.xlsx"
    output_file = path / "output" / "MYP_template_op.xlsx"
    sheet_names = ["Korea", "India", "Japan", "Thailand"]

    # Read from an Excel workbook
    wb = op.load_workbook(input_file)
    ws_list = wb.sheetnames
    print(ws_list)

    # Create an Excel workbook
    wb = op.Workbook()
    ws = wb.active

    wb.save(output_file)

    print("A file is created")


if __name__ == "__main__":
    main()
